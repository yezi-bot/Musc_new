import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


SOURCE = Path(__file__).with_name("cross_category_switch_summary.json")
OUTPUT = Path(__file__).with_name("扩展类别无RsCIN在线动态专家实验报告.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_run(run, size=10.5, bold=False, color="000000"):
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def remove_paragraph_border(paragraph_or_style):
    p_pr = paragraph_or_style._element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def style_paragraph(paragraph, size=10.5, bold=False, color="000000"):
    for run in paragraph.runs:
        style_run(run, size=size, bold=bold, color=color)


def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.35
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        style_run(lead, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        style_run(rest)
    else:
        run = paragraph.add_run(text)
        style_run(run)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    run = paragraph.add_run(text)
    style_run(run)
    return paragraph


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, (cell, text) in enumerate(zip(header.cells, headers)):
        cell.width = Cm(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "1F4E78")
        set_cell_margins(cell)
        set_cell_border(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(str(text))
        style_run(run, size=9, bold=True, color="FFFFFF")
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, values)):
            cell.width = Cm(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_shading(cell, "F2F6FA")
            set_cell_margins(cell)
            set_cell_border(cell)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 1 else WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            run = paragraph.add_run(str(value))
            style_run(run, size=8.6)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    style_run(run, size=9, color="666666")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    tail = paragraph.add_run(" 页")
    style_run(tail, size=9, color="666666")


def fmt(value):
    return "—" if value is None else f"{float(value):.4f}"


def short_variant(name):
    return "Quarantine" if name == "dynamic_dino_quarantine" else "Channel gate"


def build_document():
    runs = json.loads(SOURCE.read_text(encoding="utf-8"))["runs"]
    by_key = {(run["sequence"], run["variant"]): run for run in runs}
    sequences = [
        "bottle->capsule",
        "capsule->bottle",
        "cable->metal_nut",
        "metal_nut->cable",
    ]
    quarantine = "dynamic_dino_quarantine"
    channel_gate = "dynamic_dino_channel_gate_no_fallback"

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    add_page_number(section.footer.paragraphs[0])

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    for name, size, before, after in (
        ("Title", 22, 0, 18),
        ("Heading 1", 15, 16, 8),
        ("Heading 2", 12, 12, 6),
    ):
        style = styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = name != "Title"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        remove_paragraph_border(style)

    title = doc.add_paragraph(style="Title")
    remove_paragraph_border(title)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("扩展类别无 RsCIN 在线动态专家实验报告")
    style_run(title_run, size=22, bold=False)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    subtitle_run = subtitle.add_run("Bottle Capsule Cable Metal Nut  跨类别突变  Seed 0")
    style_run(subtitle_run, size=10.5, color="666666")

    add_body(
        doc,
        "本轮在原 Bottle 与 Cable 双向实验基础上新增 Capsule 和 Metal Nut，形成四个新的跨类别转换方向。"
        "全部实验关闭 RsCIN，图像级结果直接采用严格因果的 raw Image AUROC。结果显示，Channel gate 在四个方向上均提高整体 raw Image AUROC、Pixel AUROC 和 AUPRO，但增益大小仍具有明显方向差异。",
    )
    add_body(
        doc,
        "当前结论：新增方向支持 Channel gate 具有一定泛化性，但证据仅来自 seed 0。现在还不能据此缩短 8 张候选收集门限，应先补 seed 1 和 seed 2，确认增益和专家时延是否稳定。",
        bold_lead="当前结论：",
    )

    doc.add_heading("实验目的与设置", level=1)
    add_body(
        doc,
        "实验目标是检验当前动态专家机制从 Bottle 与 Cable 扩展到其他 MVTec AD 类别后，Channel gate 相对 Quarantine 是否仍能改善类别突变后的检测，并避免全流 RsCIN 对在线结果的影响。",
    )
    settings = [
        ("新增类别", "Capsule 和 Metal Nut"),
        ("转换方向", "Bottle→Capsule、Capsule→Bottle、Cable→Metal Nut、Metal Nut→Cable"),
        ("样本规模", "每个类别段 40 张，每条流 80 张"),
        ("随机种子", "Seed 0"),
        ("特征", "DINOv2 ViT L14，第 23 层，LNAMD r 等于 1"),
        ("方案", "Quarantine；Channel gate 无历史退回"),
        ("时间顺序", "使用步初状态评分，再更新 Channel 和专家"),
        ("RsCIN", "关闭；Image AUROC 直接等于 raw Image AUROC"),
    ]
    add_table(doc, ["项目", "设置"], settings, [4.0, 12.5])

    doc.add_heading("主要结果", level=1)
    add_body(
        doc,
        "表中的 Image AUROC 均为未经过 RsCIN 的因果 raw 指标。切换后指标只使用第二类别段。由于本轮只有一个随机种子，数值用于筛查趋势，不报告均值和标准差。",
    )
    performance_rows = []
    for sequence in sequences:
        for variant in (quarantine, channel_gate):
            run = by_key[(sequence, variant)]
            performance_rows.append(
                [
                    sequence.replace("->", "→").replace("metal_nut", "metal nut"),
                    short_variant(variant),
                    fmt(run["raw_image_auroc"]),
                    fmt(run["post_raw_image_auroc"]),
                    fmt(run["pixel_auroc"]),
                    fmt(run["aupro"]),
                ]
            )
    add_table(
        doc,
        ["方向", "方案", "Raw Image AUROC", "切换后 Raw", "Pixel AUROC", "AUPRO"],
        performance_rows,
        [3.1, 2.6, 2.8, 2.6, 2.6, 2.2],
    )

    doc.add_heading("Channel gate 相对 Quarantine 的变化", level=2)
    delta_rows = []
    delta_values = []
    for sequence in sequences:
        base = by_key[(sequence, quarantine)]
        gate = by_key[(sequence, channel_gate)]
        deltas = [
            gate["raw_image_auroc"] - base["raw_image_auroc"],
            gate["post_raw_image_auroc"] - base["post_raw_image_auroc"],
            gate["pixel_auroc"] - base["pixel_auroc"],
            gate["aupro"] - base["aupro"],
        ]
        delta_values.append(deltas)
        delta_rows.append(
            [
                sequence.replace("->", "→").replace("metal_nut", "metal nut"),
                *[f"{value:+.4f}" for value in deltas],
            ]
        )
    means = [sum(values[index] for values in delta_values) / len(delta_values) for index in range(4)]
    delta_rows.append(["四方向平均", *[f"{value:+.4f}" for value in means]])
    add_table(
        doc,
        ["方向", "Raw Image", "切换后 Raw", "Pixel AUROC", "AUPRO"],
        delta_rows,
        [4.0, 3.0, 3.0, 3.0, 3.0],
    )
    add_body(
        doc,
        "四个方向的整体 raw Image AUROC 均提高，平均增益约为 "
        f"{means[0]:.4f}；AUPRO 的平均增益约为 {means[3]:.4f}。"
        "其中 Bottle→Capsule 的提升最大，说明 Channel gate 对部分方向的切换适应更有帮助；其余方向的图像级提升较小。",
    )

    doc.add_heading("切换阶段与专家时延", level=1)
    transition_rows = []
    for sequence in sequences:
        for variant in (quarantine, channel_gate):
            run = by_key[(sequence, variant)]
            transition_rows.append(
                [
                    sequence.replace("->", "→").replace("metal_nut", "metal nut"),
                    short_variant(variant),
                    fmt(run["post10_raw_image_auroc"]),
                    fmt(run["post20_raw_image_auroc"]),
                    run["new_expert_latency"] if run["new_expert_latency"] is not None else "—",
                    f"{run['post_old_expert_reference_fraction']:.1%}",
                ]
            )
    add_table(
        doc,
        ["方向", "方案", "切换后前10张", "切换后前20张", "新专家时延", "旧专家引用占比"],
        transition_rows,
        [3.1, 2.6, 2.8, 2.8, 2.5, 2.9],
    )
    add_body(
        doc,
        "四个方向的新类别正式专家形成时延为 9 到 11 步。只有 Bottle→Capsule 在切换后仍有 22.5% 的步骤引用旧类别正式专家，其他三个方向为 0。两种方案的正式专家时间线相同，差别来自 Channel gate 在确认前使用 provisional 候选参与低权重评分。",
    )
    add_body(
        doc,
        "Bottle→Capsule 中 Channel gate 的切换后 raw AUROC 提高 0.0563，但切换后前 10 张仍然较低；Cable→Metal Nut 的前 10 张已经达到 1.0000，缩短确认门限不一定会带来收益。这说明固定把 8 张改短并不是当前证据直接支持的优化。",
    )

    doc.add_heading("运行完整性检查", level=1)
    checks = [
        "8 次正式运行和 8 份动态审计文件均已生成。",
        "所有运行都记录 use_rscin 为 false。",
        "每条运行的 image_auroc 与 raw_image_auroc 完全相同。",
        "全部运行的 causal_reference_violations 为 0。",
        "每个方向的两种方案都形成 1 个新分布正式专家；Channel gate 均完成 1 次确认。",
        "实验使用相同的 seed 0、每段 40 张、DINO 层和聚合尺度，未针对新类别重新调参。",
    ]
    for check in checks:
        add_bullet(doc, check)

    doc.add_heading("结论与下一步", level=1)
    add_body(
        doc,
        "本轮把验证范围从 Bottle 与 Cable 扩展到了 Capsule 和 Metal Nut。四个新增方向中，Channel gate 相对 Quarantine 的整体 raw Image AUROC、Pixel AUROC 和 AUPRO 均有提升，说明其效果不是只存在于 Bottle→Cable。与此同时，图像级增益从 0.0084 到 0.0280 不等，方向依赖仍然存在。",
    )
    add_body(
        doc,
        "建议下一步保持 8 张门限和其他配置不变，为这四个方向补 seed 1 和 seed 2。只有当多种子结果确认“新专家形成时延”稳定造成切换初期损失时，再做 4、6、8 张门限消融。否则应优先研究 Bottle→Capsule 中旧类别专家持续参与评分的问题。",
    )

    doc.add_heading("结果解释限制", level=1)
    limitations = [
        "本轮新增结果只有 seed 0，不能用单次运行证明统计稳定性。",
        "只扩展到 Capsule 和 Metal Nut，尚不能代表全部 MVTec AD 类别。",
        "本轮比较的是 Quarantine 与 Channel gate，不包含源码静态 fuser 和严格历史 DINO 基线。",
        "本轮没有对新增类别执行 GT only 专家污染 oracle，因此不能直接判断 provisional 污染的因果影响。",
        "关闭 RsCIN 后，报告中的图像级指标与此前使用 RsCIN 的表格不可直接混用。",
    ]
    for limitation in limitations:
        add_bullet(doc, limitation)

    doc.add_heading("可复现实验信息", level=1)
    add_body(
        doc,
        "实验入口为 MultiMuSc2 examples ablate_cross_category_switch.py。正式结果保存在当前工作树的 outputs xcat norscin s0 目录，包含逐运行 CSV、汇总 JSON 和每次运行的动态审计 JSON。代码提交 15a8153 增加了 use_rscin 开关，并在本轮跨类别入口中固定关闭。",
    )

    doc.core_properties.title = "扩展类别无 RsCIN 在线动态专家实验报告"
    doc.core_properties.subject = "跨类别动态专家扩展类别实验"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "动态专家, 跨类别, RsCIN, MVTec AD"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
